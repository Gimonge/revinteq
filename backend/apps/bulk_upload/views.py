"""
Revinteq v3 — Bulk Upload Views
Upload → Preview → Confirm → Save
"""
import logging
from django.utils import timezone
from django.db import models
from django.http import HttpResponse
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from apps.common.views import get_tenant
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser
from apps.tenants.permissions import IsClientOrAdmin
from .models import BulkSalesUpload
from .services.excel_parser import ExcelSalesParser

logger = logging.getLogger(__name__)


class BulkUploadTemplateView(APIView):
    """
    GET /api/v1/bulk-upload/template/
    Download the Excel template with correct column headers.
    No authentication required — it is a static file with no sensitive data.
    """
    permission_classes = []
    authentication_classes = []

    def get(self, request):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sales'

        # Headers
        headers = [
            'Customer Name',
            'Customer Phone',
            'Product Name *',
            'Amount (KES) *',
            'Payment Method',
            'Payment Reference',
            'Platform Source',
            'Campaign Name',
            'Sale Date (DD/MM/YYYY)',
            'Notes',
        ]
        header_fill = PatternFill(start_color='1F7A4C', end_color='1F7A4C', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, name='Calibri', size=11)
        thin = Side(style='thin', color='CCCCCC')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Example rows
        examples = [
            ['Jane Wanjiku',  '0712345678', 'Blue Summer Dress (M)', 2400, 'M-Pesa', 'QK47XY8Z21', 'Facebook', 'Summer Sale Campaign', '07/04/2026', 'Summer collection'],
            ['',              '',           'Red Handbag',           3800, 'Cash',   '',           'Instagram', '', '07/04/2026', ''],
            ['John Kamau',    '0723456789', 'White Sneakers (42)',   5500, 'Bank Deposit','TRF0012345','Organic', '', '06/04/2026','Walk-in customer'],
            ['Mary Achieng',  '0734567890', 'Gold Necklace',        12000, 'M-Pesa', 'RK92AB3C45', 'Facebook', 'Gold Collection Ad', '05/04/2026', ''],
        ]
        row_fill_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        row_font = Font(name='Calibri', size=10)

        for row_num, row_data in enumerate(examples, 2):
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.font = row_font
                cell.border = border
                cell.alignment = Alignment(vertical='center')
                if row_num % 2 == 0:
                    cell.fill = row_fill_even

        # Instructions sheet
        ws_info = wb.create_sheet(title='Instructions')
        instructions = [
            ('Revinteq Bulk Sales Upload Template', True),
            ('', False),
            ('REQUIRED COLUMNS (marked with *)', True),
            ('Product Name *  —  What was sold e.g. Blue Dress (Size M)', False),
            ('Amount (KES) *  —  Sale amount, numbers only e.g. 2400', False),
            ('', False),
            ('OPTIONAL COLUMNS', True),
            ('Customer Name  —  Customer full name e.g. Jane Wanjiku', False),
            ('Customer Phone  —  Phone number e.g. 0712345678', False),
            ('Payment Method  —  M-Pesa, Cash, Bank Deposit, Cheque, EFT, RTGS, Card, Other', False),
            ('Payment Reference  —  M-Pesa code, bank ref, cheque number etc.', False),
            ('Platform Source  —  Facebook, Instagram, Organic, Referral, Other', False),
            ('Campaign Name  —  Name of the ad campaign that brought this customer (optional)', False),
            ('Sale Date  —  Format: DD/MM/YYYY  e.g. 07/04/2026. Leave blank for today.', False),
            ('Notes  —  Any additional information', False),
            ('', False),
            ('RULES', True),
            ('• Dates cannot be in the future', False),
            ('• Dates cannot be more than 90 days in the past', False),
            ('• Amount must be greater than 0', False),
            ('• Delete these example rows before uploading', False),
            ('• Save as .xlsx format (Excel Workbook)', False),
        ]

        for row_num, (text, bold) in enumerate(instructions, 1):
            cell = ws_info.cell(row=row_num, column=1, value=text)
            cell.font = Font(name='Calibri', size=11, bold=bold,
                           color='1F7A4C' if bold else '374151')
            cell.alignment = Alignment(wrap_text=True)

        ws_info.column_dimensions['A'].width = 80

        # Column widths on main sheet
        col_widths = [22, 18, 30, 15, 20, 25, 20, 25, 25, 30]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

        ws.row_dimensions[1].height = 22

        # Write to response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="revinteq_sales_upload_template.xlsx"'
        wb.save(response)
        return response


class BulkUploadView(APIView):
    """
    POST /api/v1/bulk-upload/upload/
    Upload an Excel file. Returns a preview of valid and invalid rows.
    Does NOT save any sales yet.
    """
    permission_classes = [IsAuthenticated, IsClientOrAdmin]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response(
                {'error': True, 'message': 'No file uploaded. Please select an Excel file.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Validate file type
        filename = file.name.lower()
        if not (filename.endswith('.xlsx') or filename.endswith('.xls')):
            return Response(
                {'error': True, 'message': 'Invalid file type. Please upload an Excel file (.xlsx)'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Max 5MB
        if file.size > 5 * 1024 * 1024:
            return Response(
                {'error': True, 'message': 'File is too large. Maximum size is 5MB.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        file_bytes = file.read()

        # Parse
        parser = ExcelSalesParser(file_bytes, get_tenant(request))
        result = parser.parse()

        if not result.get('success'):
            return Response(
                {'error': True, 'message': result.get('error', 'Could not parse file.')},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Create upload session
        upload = BulkSalesUpload.objects.create(
            tenant=get_tenant(request),
            original_filename=file.name,
            total_rows=result['total_rows'],
            valid_rows=result['valid_rows'],
            error_rows=result['error_rows'],
            preview_data=result['rows'],
            status='previewed',
            uploaded_by=request.user,
        )
        # Save file
        from django.core.files.base import ContentFile
        upload.file.save(file.name, ContentFile(file_bytes), save=True)

        logger.info(
            f"Bulk upload parsed for {get_tenant(request).name}: "
            f"{result['valid_rows']}/{result['total_rows']} valid rows"
        )

        return Response({
            'upload_id': str(upload.id),
            'total_rows': result['total_rows'],
            'valid_rows': result['valid_rows'],
            'error_rows': result['error_rows'],
            'preview': result['rows'],
            'message': (
                f"Preview ready: {result['valid_rows']} valid rows, "
                f"{result['error_rows']} errors. "
                f"Review and click Confirm to save."
            )
        })


class BulkUploadConfirmView(APIView):
    """
    POST /api/v1/bulk-upload/<upload_id>/confirm/
    Confirm and save all valid rows as Sales.
    """
    permission_classes = [IsAuthenticated, IsClientOrAdmin]

    def post(self, request, upload_id):
        try:
            upload = BulkSalesUpload.objects.get(
                id=upload_id,
                tenant=get_tenant(request),
                status='previewed',
            )
        except BulkSalesUpload.DoesNotExist:
            return Response(
                {'error': True, 'message': 'Upload session not found or already processed.'},
                status=404
            )

        upload.status = 'processing'
        upload.save(update_fields=['status'])

        # Save valid rows
        from apps.sales.models import Sale
        saved = 0
        errors = []
        total_amount = 0

        for row in upload.preview_data:
            if row['status'] != 'valid':
                continue
            data = row['data']
            try:
                from datetime import date
                # Resolve campaign by name if provided
                campaign_obj = None
                campaign_name = data.get('campaign_name', '')
                if campaign_name:
                    try:
                        from apps.meta_integration.models import Campaign, AdAccount
                        tenant_obj = get_tenant(request)
                        accounts = AdAccount.objects.filter(
                            models.Q(tenant=tenant_obj) | models.Q(assigned_tenant=tenant_obj)
                        )
                        campaign_obj = Campaign.objects.filter(
                            ad_account__in=accounts,
                            name__icontains=campaign_name
                        ).first()
                    except Exception:
                        pass

                sale = Sale.objects.create(
                    tenant=get_tenant(request),
                    customer_name=data.get('customer_name', ''),
                    customer_phone=data.get('customer_phone', ''),
                    product_name=data['product_name'],
                    amount=data['amount'],
                    payment_method=data.get('payment_method', 'cash'),
                    payment_reference=data.get('payment_reference', ''),
                    platform_source=data.get('platform_source', 'organic'),
                    sale_date=date.fromisoformat(data['sale_date']),
                    notes=data.get('notes', ''),
                    campaign=campaign_obj,
                    is_confirmed=True,
                    bulk_upload=upload,
                    created_by=request.user,
                )
                # Auto-create/update customer record if name or phone provided
                try:
                    from apps.customers.models import Customer
                    Customer.get_or_create_from_sale(sale)
                except Exception:
                    pass
                saved += 1
                total_amount += float(data['amount'])
            except Exception as e:
                errors.append({'row': row['row_num'], 'error': str(e)})

        upload.saved_rows = saved
        upload.error_summary = errors
        upload.status = 'completed'
        upload.completed_at = timezone.now()
        upload.save()

        # Invalidate dashboard cache
        from apps.metrics.tasks import invalidate_dashboard_cache
        invalidate_dashboard_cache.delay(str(get_tenant(request).id))

        logger.info(
            f"Bulk upload confirmed: {saved} sales saved for {get_tenant(request).name}. "
            f"Total: KES {total_amount:,.2f}"
        )

        return Response({
            'message': f'Successfully saved {saved} sales.',
            'saved_rows': saved,
            'total_amount': total_amount,
            'errors': errors,
        })
