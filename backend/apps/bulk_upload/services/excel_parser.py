"""
Revinteq v3 — Excel Parser
Uses openpyxl only — no pandas, no C++ build tools needed on Windows.
"""
import io
import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

logger = logging.getLogger(__name__)

PAYMENT_METHOD_MAP = {
    'mpesa':          'mpesa_manual',
    'm-pesa':         'mpesa_manual',
    'm pesa':         'mpesa_manual',
    'mpesa manual':   'mpesa_manual',
    'cash':           'cash',
    'bank':           'bank_deposit',
    'bank deposit':   'bank_deposit',
    'eft':            'eft',
    'rtgs':           'rtgs',
    'standing order': 'standing_order',
    'cheque':         'cheque',
    'check':          'cheque',
    'card':           'card',
    'other':          'other',
}

PLATFORM_MAP = {
    'facebook':  'facebook',
    'fb':        'facebook',
    'instagram': 'instagram',
    'ig':        'instagram',
    'messenger': 'messenger',
    'whatsapp':  'facebook',
    'organic':   'organic',
    'walk-in':   'organic',
    'walkin':    'organic',
    'referral':  'organic',
    'other':     'organic',
}


class ExcelSalesParser:
    """
    Parses an Excel file uploaded for bulk sales import.
    Accepts bytes (file.read()) or a file path string.

    Usage:
        parser = ExcelSalesParser(file_bytes, tenant)
        result = parser.parse()
        # result['success'] bool
        # result['rows']    list of {row_num, status, data, errors}
        # result['valid_rows'] int
        # result['error_rows'] int
        # result['total_rows'] int
    """

    # Map from display column headers to internal keys
    COLUMN_ALIASES = {
        'product name *':        'product_name',
        'product name':          'product_name',
        'product':               'product_name',
        'amount (kes) *':        'amount',
        'amount *':              'amount',
        'amount':                'amount',
        'payment method':        'payment_method',
        'payment':               'payment_method',
        'payment reference':     'payment_reference',
        'reference':             'payment_reference',
        'platform source':       'platform_source',
        'platform':              'platform_source',
        'source':                'platform_source',
        'sale date (dd/mm/yyyy)':'sale_date',
        'sale date':             'sale_date',
        'date':                  'sale_date',
        'notes':                 'notes',
        'campaign name':         'campaign_name',
        'campaign':              'campaign_name',
        'customer name':         'customer_name',
        'customer':              'customer_name',
        'customer phone':        'customer_phone',
        'phone':                 'customer_phone',
        'phone number':          'customer_phone',
        'mobile':                'customer_phone',
    }

    def __init__(self, file_data, tenant=None):
        """
        file_data: bytes from file.read(), or a string file path
        tenant:    optional tenant object (for context)
        """
        self.file_data = file_data
        self.tenant = tenant

    def parse(self) -> dict:
        try:
            import openpyxl
        except ImportError:
            return {'success': False, 'error': 'openpyxl is not installed. Run: pip install openpyxl'}

        try:
            if isinstance(self.file_data, (str, bytes)) and not isinstance(self.file_data, bytes):
                # File path
                wb = openpyxl.load_workbook(self.file_data, read_only=True, data_only=True)
            else:
                # Bytes
                wb = openpyxl.load_workbook(
                    io.BytesIO(self.file_data), read_only=True, data_only=True
                )
            ws = wb.active
        except Exception as e:
            return {'success': False, 'error': f'Could not open file: {e}'}

        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not all_rows:
            return {'success': False, 'error': 'The file is empty.'}

        # Build column map from header row
        raw_headers = [
            str(h).strip().lower() if h is not None else ''
            for h in all_rows[0]
        ]
        col_map = {}
        for idx, raw_h in enumerate(raw_headers):
            internal = self.COLUMN_ALIASES.get(raw_h)
            if internal:
                col_map[internal] = idx

        if 'product_name' not in col_map:
            return {
                'success': False,
                'error': (
                    'Could not find required columns. '
                    'Please download the template and use it as your starting point.'
                )
            }

        rows_output = []
        valid_count = 0
        error_count = 0

        for row_idx, raw_row in enumerate(all_rows[1:], start=2):
            # Skip completely empty rows
            if all(v is None or str(v).strip() == '' for v in raw_row):
                continue

            row_errors = []
            data = {}

            def get(key, default=''):
                idx = col_map.get(key)
                if idx is None or idx >= len(raw_row):
                    return default
                val = raw_row[idx]
                return val if val is not None else default

            # Product name (required)
            product = str(get('product_name', '')).strip()
            if not product or product.lower() in ('none', 'nan'):
                row_errors.append('Product name is required.')
            else:
                data['product_name'] = product

            # Amount (required)
            raw_amount = get('amount', '')
            try:
                clean = str(raw_amount).replace(',', '').replace('KES', '').replace(' ', '').strip()
                amount = Decimal(clean)
                if amount <= 0:
                    row_errors.append('Amount must be greater than 0.')
                else:
                    data['amount'] = str(amount)
            except (InvalidOperation, ValueError):
                row_errors.append(f"Invalid amount: '{raw_amount}'. Use numbers only e.g. 2400")

            # Payment method (optional, default cash)
            raw_pm = str(get('payment_method', 'cash')).strip().lower()
            pm = PAYMENT_METHOD_MAP.get(raw_pm, 'cash')
            data['payment_method'] = pm

            # Payment reference (optional)
            ref = str(get('payment_reference', '')).strip()
            if ref and ref.lower() not in ('none', 'nan'):
                data['payment_reference'] = ref

            # Platform (optional, default facebook)
            raw_plat = str(get('platform_source', 'facebook')).strip().lower()
            data['platform_source'] = PLATFORM_MAP.get(raw_plat, 'facebook')

            # Sale date (optional, default today)
            raw_date = get('sale_date', '')
            if raw_date and str(raw_date).strip() not in ('', 'none', 'nan'):
                parsed = _parse_date(raw_date)
                if parsed is None:
                    row_errors.append(
                        f"Invalid date '{raw_date}'. Use DD/MM/YYYY e.g. 07/04/2026"
                    )
                elif parsed > date.today():
                    row_errors.append('Sale date cannot be in the future.')
                else:
                    data['sale_date'] = parsed.isoformat()
            else:
                data['sale_date'] = date.today().isoformat()

            # Customer name (optional)
            cust_name = str(get('customer_name', '')).strip()
            if cust_name and cust_name.lower() not in ('none', 'nan'):
                data['customer_name'] = cust_name

            # Customer phone (optional)
            cust_phone = str(get('customer_phone', '')).strip()
            if cust_phone and cust_phone.lower() not in ('none', 'nan'):
                data['customer_phone'] = cust_phone

            # Campaign name (optional — for attribution)
            campaign_name = str(get('campaign_name', '')).strip()
            if campaign_name and campaign_name.lower() not in ('none', 'nan'):
                data['campaign_name'] = campaign_name

            # Notes (optional)
            notes = str(get('notes', '')).strip()
            if notes and notes.lower() not in ('none', 'nan'):
                data['notes'] = notes

            if row_errors:
                error_count += 1
                rows_output.append({
                    'row_num': row_idx,
                    'status':  'error',
                    'data':    data if data else None,
                    'errors':  row_errors,
                })
            else:
                valid_count += 1
                rows_output.append({
                    'row_num': row_idx,
                    'status':  'valid',
                    'data':    data,
                    'errors':  [],
                })

        return {
            'success':    True,
            'total_rows': valid_count + error_count,
            'valid_rows': valid_count,
            'error_rows': error_count,
            'rows':       rows_output,
        }


def _parse_date(raw) -> date | None:
    """Try multiple date formats. Returns date or None."""
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw

    s = str(raw).strip()
    if not s or s.lower() in ('none', 'nan'):
        return None

    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y',
                '%d %b %Y', '%d %B %Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None
